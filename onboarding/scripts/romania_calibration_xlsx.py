# -*- coding: utf-8 -*-
"""Living calibration workbook: one column per run, three metric sheets.

calibration_tables.xlsx (disrupt-sc-data/Romania) holds the calibration
tables that used to be rebuilt into the validation note docx:

  Runs        run id | date | label            (one row per update)
  Economic    model totals vs MRIO
  ModalSplit  per-cargo mode shares + tkm levels vs Eurostat 2023
  BCP         one row per border-crossing point and mode vs the WB counts

Design: VALUES ONLY, no formulas - this script is the calculator (it reads
the run's own exports: firm_data, trade_data, transport_edges_with_flows).
Static columns (labels, targets, sources, the manual Verdict column on BCP)
are written once at creation and never touched again; each invocation adds
or refreshes one run column, so the workbook accumulates the calibration
lineage side by side.

Usage (dsc env):
    python romania_calibration_xlsx.py <run_id> [--label "short description"]
"""

from __future__ import annotations

import argparse
import datetime as dt
import math
import sys
from pathlib import Path

import geopandas as gpd
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

REPO = Path(r"C:\Users\Celian\OneDrive\DisruptSC\disrupt-sc")
DATA = Path(r"C:\Users\Celian\OneDrive\DisruptSC\disrupt-sc-data\Romania")
XLSX = DATA / "calibration_tables.xlsx"

ARIAL = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HDR_FILL = PatternFill("solid", fgColor="DCE6F1")
EDIT_FILL = PatternFill("solid", fgColor="FFF2CC")  # manual columns

ECONOMIC_ROWS = [
    ("Gross output, % of MRIO", "100", "model firm production x52 vs MRIO output (598.3 bUSD)", None),
    ("Sectors within +/-15% of MRIO output (of 50)", "50", "per-sector production vs mrio_by_sector.csv", None),
    ("Imports (countries' sales into scope), % of MRIO", "100", "vs MRIO 134,615 mUSD/yr; INCLUDES transit throughput since v8", None),
    ("Exports (purchases from scope), % of MRIO", "100", "vs MRIO 115,146 mUSD/yr; INCLUDES transit throughput since v8", None),
]

MODAL_ROWS = [
    ("Dry bulk - road, % tkm", "35.4", "Eurostat 2023 NST->cargo mapping; MODEL BASIS: whole network incl. foreign legs (canonical since v5)", None),
    ("Dry bulk - rail, % tkm", "22.9", "", None),
    ("Dry bulk - IWW, % tkm", "41.7", "", None),
    ("Container - road, % tkm", "89.2", "", None),
    ("Container - rail, % tkm", "6.5", "", None),
    ("Container - IWW, % tkm", "4.3", "", None),
    ("Liquid bulk - road, % tkm", "21.6", "", None),
    ("Liquid bulk - rail, % tkm", "68.5", "", None),
    ("Liquid bulk - IWW, % tkm", "10.0", "", None),
    ("All inland - road, % tkm", "72.6", "MODEL BASIS: domestic territorial. Share target caveat: Eurostat road tkm = registered hauliers anywhere", None),
    ("All inland - rail, % tkm", "14.1", "", None),
    ("All inland - IWW, % tkm", "13.3", "", None),
    ("Rail, bn tkm/yr", "12.65", "Eurostat 2023 level", None),
    ("IWW, bn tkm/yr", "11.96", "Eurostat 2023 level; model includes UA transit riding the Danube", None),
    ("Road, bn tkm/yr", "(65.2)", "NOT comparable: registered-haulier basis", None),
]

# (row label, WB data, note, verdict seed, extractor key)
BCP_ROWS = [
    ("HU - Nadlac/Curtici (road)", "-", "no WB data for EU crossings", "dominant EU road gate; winner-take-all concentration", ("gaz", "Nadlac", "roads")),
    ("HU - Bors/Episcopia (road)", "-", "", "secondary NW gate", ("gaz", "Bors", "roads")),
    ("HU - Bors/Episcopia (rail)", "-", "", "main EU rail crossing", ("gaz", "Bors", "railways")),
    ("HU - Petea (road)", "-", "", "northern gate", ("gaz", "Petea", "roads")),
    ("BG - Giurgiu-Ruse (road)", "-", "", "main southern gate (TR + BG + transit)", ("gaz", "Giurgiu", "roads")),
    ("BG - Calafat-Vidin (road)", "-", "", "minor", ("gaz", "Calafat", "roads")),
    ("BG - Calafat-Vidin (rail)", "-", "", "minor", ("gaz", "Calafat", "railways")),
    ("RS - Iron Gates locks (barge)", "~5.5 (viadonau)", "non-WB anchor: lock statistics", "over: upstream transit + import bulk", ("gaz", "IronGates", "waterways")),
    ("UA - Siret/Porubne (road)", "1.4*", "* UA-exit direction only; model = both directions", "under; corridor partly rides rail/barge", ("name", "Porubne-Siret", None)),
    ("UA - Vadul-Siret/Vicsani (rail)", "2.0*", "", "under x10: central-UA grain sits in the Kyiv sub-bloc", ("name", "gauge break Vadul-Siret", None)),
    ("UA - Halmeu/Dyakove (road)", "0.45*", "", "under: Kyiv-centric geometry bypasses Transcarpathia", ("name", "Dyakove-Halmeu", None)),
    ("UA - Halmeu/Dyakove (rail)", "0.44*", "", "idem", ("name", "gauge break Dyakove/Halmeu", None)),
    ("UA - Isaccea/Orlivka (ferry)", "0.53*", "", "Odesa-Constanta block sits on a barge/ferry knife edge", ("name", "Orlivka-Isaccea", None)),
    ("UA - Danube ports, barge (Reni+Izmail)", "6-8", "2024 Danube-port band (WB report + port statistics)", "in band", ("name2", "Izmail-Danube", "Reni-Danube")),
    ("MD - Albita/Leuseni (road)", "2.5-4.0", "27k trucks/month both directions incl. empties", "low side (counts include empties and MD-EU transit)", ("name", "Leuseni-Albita", None)),
    ("MD - Sculeni (road)", "~1.3", "", "under: winner-take-all sends Chisinau road flow to Albita", ("name", "UAMD:Sculeni", None)),
    ("MD - Oancea/Cahul (road)", "~1.2", "", "under: idem", ("name", "Cahul-Oancea", None)),
    ("MD - Giurgiulesti-Galati (road)", "~1.3", "", "under: road flow displaced by the cheap 1520 mm rail", ("name", "Giurgiulesti-Galati", None)),
    ("MD - Ungheni (rail)", "0.3-0.5", "10.5k wagons/yr both directions at ~50 t", "on target", ("name", "stitch railways @(27.81,47.23)", None)),
    ("MD - Giurgiulesti CFR (rail)", "0.5-0.9", "18.3k wagons/yr", "over: southern line overshoots (mirror of Vadul deficit)", ("name", "stitch railways @(28.20,45.47)", None)),
]

GAZ = {"Nadlac": (20.9, 46.2), "Bors": (21.9, 47.1), "Petea": (23.1, 47.9),
       "Giurgiu": (25.9, 43.8), "Calafat": (22.9, 44.0), "IronGates": (20.4, 45.0),
       "Ungheni": (27.8, 47.2), "GiurgiulestiR": (28.2, 45.5)}


def hav(a, b):
    lo1, la1, lo2, la2 = map(math.radians, [a[0], a[1], b[0], b[1]])
    return 6371 * 2 * math.asin(math.sqrt(
        math.sin((la2 - la1) / 2) ** 2
        + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2))


def compute(run_id: str) -> dict:
    run = REPO / "output" / "Romania" / run_id
    out: dict = {}

    fd = pd.read_csv(run / "firm_data.csv")
    fd = fd[fd["time_step"] == 0]
    ms = pd.read_csv(run / "mrio_by_sector.csv").set_index("sector")
    model_out = fd["production"].sum() * 52
    out["Gross output, % of MRIO"] = round(100 * model_out / ms["mrio_output"].sum(), 1)
    per = fd.groupby("sector")["production"].sum() * 52
    ratio = (per / ms["mrio_output"]).dropna()
    out["Sectors within +/-15% of MRIO output (of 50)"] = int(((ratio - 1).abs() <= 0.15).sum())

    td = pd.read_csv(run / "trade_data.csv")
    td = td[(td["time_step"] == 0) & (td["country"] != "ROU")]  # ROU row mirrors scope totals
    out["Imports (countries' sales into scope), % of MRIO"] = round(100 * td["exports_value"].sum() * 52 / 134615, 1)
    out["Exports (purchases from scope), % of MRIO"] = round(100 * td["imports_value"].sum() * 52 / 115146, 1)

    e = gpd.read_file(run / "transport_edges_with_flows_0.geojson")
    e["tons"] = e["flow_total_tons"].fillna(0)
    dom = e[(e["foreign"].fillna(0) != 1)]
    inland = dom[dom["type"].isin(["roads", "railways", "waterways"])]
    mode_map = {"roads": "road", "railways": "rail", "waterways": "IWW"}
    # per-cargo shares over the WHOLE network incl. foreign legs - the basis
    # eurostat_mode_targets.py --compare has used since v5 (comparability);
    # the All-inland rows and tkm levels below are DOMESTIC territorial.
    whole = e[e["type"].isin(["roads", "railways", "waterways"])]
    for label, ct in (("Dry bulk", "tons_dry_bulk"), ("Container", "tons_container"),
                      ("Liquid bulk", "tons_liquid_bulk")):
        tkm = whole.groupby("type").apply(
            lambda g, c=ct: (g[c].fillna(0) * g["km"]).sum(), include_groups=False)
        tot = tkm.sum() or 1.0
        for mode, short in mode_map.items():
            out[f"{label} - {short}, % tkm"] = round(100 * tkm.get(mode, 0.0) / tot, 1)
    tkm_all = inland.groupby("type").apply(
        lambda g: (g["tons"] * g["km"]).sum(), include_groups=False)
    tot = tkm_all.sum() or 1.0
    for mode, short in mode_map.items():
        out[f"All inland - {short}, % tkm"] = round(100 * tkm_all.get(mode, 0.0) / tot, 1)
    out["Rail, bn tkm/yr"] = round(tkm_all.get("railways", 0.0) * 52 / 1e9, 1)
    out["IWW, bn tkm/yr"] = round(tkm_all.get("waterways", 0.0) * 52 / 1e9, 1)
    out["Road, bn tkm/yr"] = round(tkm_all.get("roads", 0.0) * 52 / 1e9, 1)

    # BCP: gazetteer over stitch/bridge edges, plus name-marked crossings
    st = e[(e["class"].astype(str) == "stitch")
           | e["name"].astype(str).str.contains("bridge:", na=False)]
    cent = st.geometry.centroid
    gaz_mt: dict = {}
    for i, r in st.iterrows():
        c = (cent[i].x, cent[i].y)
        lab = min(GAZ, key=lambda k: hav(GAZ[k], c))
        key = (lab, r["type"])
        gaz_mt[key] = gaz_mt.get(key, 0.0) + r["tons"] * 52 / 1e6
    names = e["name"].astype(str)
    for label, wb, note, verdict, (kind, a, b) in BCP_ROWS:
        if kind == "gaz":
            out[label] = round(gaz_mt.get((a, b), 0.0), 2)
        elif kind == "name":
            out[label] = round(e.loc[names.str.contains(a, regex=False), "tons"].max() * 52 / 1e6, 2)
        elif kind == "name2":
            v = sum(e.loc[names.str.contains(t, regex=False), "tons"].sum() for t in (a, b))
            out[label] = round(v * 52 / 1e6, 2)
    return out


def style_row(ws, row, font=ARIAL):
    for cell in ws[row]:
        cell.font = font


def create_workbook():
    wb = Workbook()
    runs = wb.active
    runs.title = "Runs"
    runs.append(["run_id", "date", "label"])
    style_row(runs, 1, BOLD)
    runs.append(["", "", "Legend: this workbook is written by onboarding/scripts/"
                 "romania_calibration_xlsx.py - values only, no formulas (the script is the "
                 "calculator). Editable by hand: this label column, and the yellow Verdict "
                 "column on BCP. Each run adds one column to the three metric sheets."])
    style_row(runs, 2)
    for cell in runs[1]:
        cell.fill = HDR_FILL
    runs.column_dimensions["A"].width = 18
    runs.column_dimensions["B"].width = 12
    runs.column_dimensions["C"].width = 110

    for sheet, rows, has_verdict in (("Economic", ECONOMIC_ROWS, False),
                                     ("ModalSplit", MODAL_ROWS, False),
                                     ("BCP", BCP_ROWS, True)):
        ws = wb.create_sheet(sheet)
        hdr = ["Metric" if sheet != "BCP" else "Border crossing",
               "Target" if sheet != "BCP" else "WB data", "Source / notes"]
        if has_verdict:
            hdr.append("Verdict (manual)")
        ws.append(hdr)
        for cell in ws[1]:
            cell.font = BOLD
            cell.fill = HDR_FILL
        for row in rows:
            label, target, note = row[0], row[1], row[2]
            line = [label, target, note]
            if has_verdict:
                line.append(row[3])
            ws.append(line)
            style_row(ws, ws.max_row)
            if has_verdict:
                ws.cell(ws.max_row, 4).fill = EDIT_FILL
        widths = [42, 14, 55] + ([46] if has_verdict else [])
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = w
        ws.freeze_panes = ws.cell(2, len(hdr) + 1)
    return wb


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("run_id")
    ap.add_argument("--label", default="")
    args = ap.parse_args()

    values = compute(args.run_id)

    if XLSX.exists():
        wb = load_workbook(XLSX)
    else:
        wb = create_workbook()
        print(f"created {XLSX.name}")

    runs = wb["Runs"]
    existing = [runs.cell(r, 1).value for r in range(2, runs.max_row + 1)]
    if args.run_id not in existing:
        runs.append([args.run_id, dt.date.today().isoformat(), args.label])
        style_row(runs, runs.max_row)
    elif args.label:
        runs.cell(existing.index(args.run_id) + 2, 3, args.label)

    for sheet in ("Economic", "ModalSplit", "BCP"):
        ws = wb[sheet]
        hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = hdr.index(args.run_id) + 1 if args.run_id in hdr else ws.max_column + 1
        ws.cell(1, col, args.run_id).font = BOLD
        ws.cell(1, col).fill = HDR_FILL
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 16
        for r in range(2, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if label in values:
                cell = ws.cell(r, col, values[label])
                cell.font = ARIAL
                cell.number_format = "0.0" if sheet != "BCP" else "0.00"

    wb.save(XLSX)
    print(f"updated {XLSX.name} with run {args.run_id} "
          f"({sum(1 for k in values)} metrics)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
